from __future__ import annotations

import io
import openpyxl

from detect_ai.modules.reporting.public import (
    build_class_conduct_rows,
    calculate_all_classes_exam_ban,
    calculate_class_exam_ban,
    generate_excel_exam_ban,
    summarize_conduct,
)
from detect_ai.modules.reporting.application.build import (
    EXAM_BAN_HEADERS,
    calculate_all_classes_exam_ban as app_all,
    calculate_class_exam_ban as app_class,
    generate_excel_exam_ban as app_excel,
)
from detect_ai.modules.reporting.presentation.routes import router as admin_reports_router
from detect_ai.modules.reporting.presentation.portal_routes import router as student_portal_router
from detect_ai.modules.reporting.application import build as report_service
from detect_ai.modules.reporting.presentation import routes as routes_admin_reports
from detect_ai.modules.reporting.presentation import portal_routes as routes_student_portal


def test_public_reexports():
    assert calculate_class_exam_ban is app_class
    assert calculate_all_classes_exam_ban is app_all
    assert generate_excel_exam_ban is app_excel


def test_generate_excel_exam_ban():
    data = [
        {
            "student_id": 1,
            "student_code": "SV001",
            "full_name": "Nguyen Van A",
            "class_name": "D20CNTT1",
            "course_name": "Python",
            "absent_count": 3,
            "half_absent_count": 1,
            "rate_semester": 23.3,
            "status_label": "BỊ CẤM THI",
            "is_banned": True,
            "is_warning": False,
        },
        {
            "student_id": 2,
            "student_code": "SV002",
            "full_name": "Tran Thi B",
            "class_name": "D20CNTT1",
            "course_name": "Python",
            "absent_count": 1,
            "half_absent_count": 0,
            "rate_semester": 6.7,
            "status_label": "ĐỦ ĐIỀU KIỆN",
            "is_banned": False,
            "is_warning": False,
        },
    ]
    raw_bytes = generate_excel_exam_ban(data, title="BÁO CÁO CẤM THI", subtitle="Lớp D20CNTT1")
    assert isinstance(raw_bytes, bytes)
    assert len(raw_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    ws = wb.active
    assert ws.title == "Danh sách cấm thi"

    # Row 1: Title
    assert ws["A1"].value == "BÁO CÁO CẤM THI"
    # Row 2: Subtitle
    assert ws["A2"].value == "Lớp D20CNTT1"
    # Row 4: Headers
    headers = [ws.cell(row=4, column=col).value for col in range(1, len(EXAM_BAN_HEADERS) + 1)]
    assert headers == EXAM_BAN_HEADERS

    # Row 5: First student
    assert ws.cell(row=5, column=2).value == "SV001"
    assert ws.cell(row=5, column=3).value == "Nguyen Van A"
    assert ws.cell(row=5, column=9).value == "BỊ CẤM THI"

    # Row 6: Second student
    assert ws.cell(row=6, column=2).value == "SV002"
    assert ws.cell(row=6, column=9).value == "ĐỦ ĐIỀU KIỆN"


def test_routers_structure():
    assert admin_reports_router.prefix == "/api/v1/admin/reports"
    admin_paths = [r.path for r in admin_reports_router.routes]
    assert "/api/v1/admin/reports/export-exam-ban-all-excel" in admin_paths

    assert student_portal_router.prefix == "/api/v1/student"
    portal_paths = [r.path for r in student_portal_router.routes]
    assert "/api/v1/student/schedule" in portal_paths
    assert "/api/v1/student/attendance" in portal_paths
    assert "/api/v1/student/conduct-ledger" in portal_paths


def test_shims_completeness():
    assert report_service.calculate_class_exam_ban is app_class
    assert report_service.calculate_all_classes_exam_ban is app_all
    assert report_service.generate_excel_exam_ban is app_excel
    assert routes_admin_reports.router is admin_reports_router
    assert routes_student_portal.router is student_portal_router


class _Profile:
    def __init__(self, total_point, class_name="D20CNTT1"):
        self.total_point = total_point
        self.class_name = class_name


class _User:
    def __init__(self, id, user_code, full_name, total_point, is_officer=False):
        self.id = id
        self.user_code = user_code
        self.full_name = full_name
        self.student_profile = _Profile(total_point)
        self.student_profile.is_officer = is_officer


class _Sched:
    def __init__(self, id):
        self.id = id


class _Entry:
    def __init__(self, student_id, schedule_id, point_delta):
        self.student_id = student_id
        self.schedule_id = schedule_id
        self.point_delta = point_delta


def test_build_class_conduct_rows():
    students = [
        _User(1, "SV001", "Nguyen Van A", 95),
        _User(2, "SV002", "Tran Thi B", 40),
        _User(3, "SV003", "Le Van C", 95, is_officer=True),
    ]
    schedules = [_Sched(10), _Sched(11)]
    ledger = [
        _Entry(1, 10, 2),
        _Entry(1, 11, -1),
        _Entry(2, 10, -5),
    ]
    rows = build_class_conduct_rows(students, schedules, ledger)
    by_id = {r["student_id"]: r for r in rows}
    a, c, b = by_id[1], by_id[3], by_id[2]
    assert (a["total_point"], a["conduct_rank"], a["class_delta"], a["class_entries"]) == (95, "Tốt", 1, 2)
    assert (c["total_point"], c["conduct_rank"], c["is_officer"]) == (95, "Xuất sắc", True)
    assert (b["total_point"], b["conduct_rank"], b["class_delta"], b["class_entries"]) == (40, "Yếu", -5, 1)
    assert a["evaluated_sessions"] == 2
    summary = summarize_conduct(rows)
    assert summary == {"total": 3, "by_rank": {"Tốt": 1, "Xuất sắc": 1, "Yếu": 1}}


def test_parse_month_param():
    from detect_ai.modules.reporting.public import parse_month_param

    assert parse_month_param(None) is None
    assert parse_month_param("") is None
    assert parse_month_param("2026-10") == (2026, 10)
    for bad in ["2026", "10-2026", "2026-13", "2026-00", "abcd-ef", "2026-1"]:
        if bad == "2026-1":
            assert parse_month_param(bad) == (2026, 1)
            continue
        try:
            parse_month_param(bad)
        except ValueError:
            continue
        raise AssertionError(f"expected ValueError for {bad!r}")


def test_mailer_config_from_db_map():
    from detect_ai.shared.notifications.mailer import config_from_map

    cfg = config_from_map({
        "SMTP_ENABLED": "true",
        "SMTP_HOST": "smtp.gmail.com",
        "SMTP_PORT": "587",
        "SMTP_USER": "a@b.c",
        "SMTP_PASSWORD": "secret",
        "SMTP_FROM": "a@b.c",
        "SMTP_USE_TLS": "true",
        "SMTP_TIMEOUT_S": "10",
    })
    assert cfg.configured is True
    assert (cfg.host, cfg.port, cfg.user, cfg.password) == ("smtp.gmail.com", 587, "a@b.c", "secret")
    assert config_from_map({}).configured is False
    assert config_from_map({"SMTP_ENABLED": "true"}).configured is False


def test_smtp_settings_validation():
    import pytest
    from fastapi import HTTPException
    from detect_ai.modules.reporting.presentation.settings_routes import _validate_setting_value

    assert _validate_setting_value("SMTP_HOST", "smtp.gmail.com") == "smtp.gmail.com"
    assert _validate_setting_value("SMTP_ENABLED", "true") == "true"
    assert _validate_setting_value("SMTP_PORT", "587") == "587"
    with pytest.raises(HTTPException):
        _validate_setting_value("SMTP_PORT", "abc")
    with pytest.raises(HTTPException):
        _validate_setting_value("SMTP_PORT", "0")
    with pytest.raises(HTTPException):
        _validate_setting_value("SMTP_ENABLED", "yes-please")


def test_generate_excel_conduct():
    import io
    import openpyxl
    from detect_ai.modules.reporting.public import CONDUCT_EXCEL_HEADERS, generate_excel_conduct

    rows = [
        {"student_code": "SV001", "full_name": "Nguyen Van A", "class_name": "D20CNTT1",
         "total_point": 95, "conduct_rank": "Tốt", "class_delta": 1, "class_entries": 2},
        {"student_code": "SV002", "full_name": "Tran Thi B", "class_name": "D20CNTT1",
         "total_point": 40, "conduct_rank": "Yếu", "class_delta": -5, "class_entries": 1},
    ]
    raw = generate_excel_conduct(rows, title="ĐIỂM RÈN LUYỆN LỚP CS101", subtitle="Môn học: Python")
    assert isinstance(raw, bytes) and len(raw) > 0
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    assert ws.title == "Điểm rèn luyện"
    headers = [ws.cell(row=4, column=col).value for col in range(1, len(CONDUCT_EXCEL_HEADERS) + 1)]
    assert headers == CONDUCT_EXCEL_HEADERS
    assert ws.cell(row=5, column=2).value == "SV001"
    assert ws.cell(row=5, column=5).value == 95
    assert ws.cell(row=5, column=7).value == "+1"
    assert ws.cell(row=6, column=7).value == "-5"


def test_conduct_excel_routes_registered():
    from detect_ai.modules.attendance.public import teacher_router

    teacher_paths = [r.path for r in teacher_router.routes]
    assert "/api/v1/teacher/classes/{class_id}/export-conduct-excel" in teacher_paths
    admin_paths = [r.path for r in admin_reports_router.routes]
    assert "/api/v1/admin/reports/classes/{class_id}/export-conduct-excel" in admin_paths


def test_conduct_routes_registered():
    from detect_ai.modules.attendance.public import teacher_router

    teacher_paths = [r.path for r in teacher_router.routes]
    assert "/api/v1/teacher/classes/{class_id}/conduct-points" in teacher_paths
    admin_paths = [r.path for r in admin_reports_router.routes]
    assert "/api/v1/admin/reports/classes/{class_id}/conduct-points" in admin_paths
