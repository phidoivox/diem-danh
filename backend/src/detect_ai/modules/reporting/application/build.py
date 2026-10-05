from __future__ import annotations

import io
from collections import Counter
from datetime import date, datetime
from typing import Any, Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from detect_ai.shared.kernel.timeutil import now_vn, today_vn
from detect_ai.models.academic import Class, ClassStudent, Schedule
from detect_ai.models.attendance_v2 import Attendance, AttendanceStatus
from detect_ai.models.user import User
from detect_ai.shared.kernel.settings import load_exam_ban_thresholds
# Vắng quy đổi: V = N_vắng + 0.5 × N_vắng_nửa_buổi (EXCUSED không tính).
# Hằng số dùng chung từ domain/policies/exam_ban_policy (nguồn duy nhất).
from detect_ai.domain.policies.exam_ban_policy import HALF_ABSENT_WEIGHT  # noqa: E402

EXAM_BAN_HEADERS = [
    "STT",
    "Mã Sinh Viên",
    "Họ và Tên",
    "Lớp Sinh Hoạt",
    "Môn Học",
    "Số Buổi Vắng",
    "Số Buổi Vắng Nửa Buổi",
    "Tỷ Lệ Vắng Cả Kỳ (%)",
    "Trạng Thái",
]

NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
WHITE_FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="4B5563")
BANNED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
BANNED_FONT = Font(name="Calibri", size=11, bold=True, color="991B1B")
WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WARNING_FONT = Font(name="Calibri", size=11, bold=True, color="92400E")
REGULAR_FONT = Font(name="Calibri", size=11)
_THIN_SIDE = Side(border_style="thin", color="D1D5DB")
CELL_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def _ensure_styles(wb: Workbook) -> None:
    """Đăng ký NamedStyle dùng chung cho báo cáo cấm thi (idempotent).

    Kiểm tra `wb.named_styles` trước khi thêm để không trùng tên khi hàm
    generate được gọi nhiều lần trên cùng workbook.
    """
    existing = {getattr(s, "name", s) for s in wb.named_styles}

    def _add(name: str, **kwargs: Any) -> None:
        if name not in existing:
            wb.add_named_style(NamedStyle(name=name, **kwargs))
            existing.add(name)

    _add("exam_title", font=TITLE_FONT, alignment=CENTER)
    _add("exam_subtitle", font=SUBTITLE_FONT, alignment=CENTER)
    _add(
        "exam_header",
        fill=NAVY_FILL,
        font=WHITE_FONT_BOLD,
        alignment=CENTER_WRAP,
        border=CELL_BORDER,
    )
    _add("exam_data", font=REGULAR_FONT, alignment=CENTER, border=CELL_BORDER)
    _add("exam_data_left", font=REGULAR_FONT, alignment=LEFT, border=CELL_BORDER)
    _add(
        "exam_banned",
        fill=BANNED_FILL,
        font=REGULAR_FONT,
        alignment=CENTER,
        border=CELL_BORDER,
    )
    _add(
        "exam_banned_left",
        fill=BANNED_FILL,
        font=REGULAR_FONT,
        alignment=LEFT,
        border=CELL_BORDER,
    )
    _add(
        "exam_banned_status",
        fill=BANNED_FILL,
        font=BANNED_FONT,
        alignment=CENTER,
        border=CELL_BORDER,
    )
    _add(
        "exam_warning",
        fill=WARNING_FILL,
        font=REGULAR_FONT,
        alignment=CENTER,
        border=CELL_BORDER,
    )
    _add(
        "exam_warning_left",
        fill=WARNING_FILL,
        font=REGULAR_FONT,
        alignment=LEFT,
        border=CELL_BORDER,
    )
    _add(
        "exam_warning_status",
        fill=WARNING_FILL,
        font=WARNING_FONT,
        alignment=CENTER,
        border=CELL_BORDER,
    )


async def calculate_class_exam_ban(
    session: AsyncSession,
    class_id: int,
    reference_date: Optional[date] = None,
    thresholds: Optional[tuple[float, float]] = None,
) -> Optional[Dict[str, Any]]:
    """Calculate exam-ban metrics and status for all students in a specific class.

    `thresholds` = (cảnh báo %, cấm thi %) đọc từ system_settings. Truyền vào khi gọi
    hàng loạt để khỏi truy vấn cấu hình lặp lại cho từng lớp.
    """
    if reference_date is None:
        reference_date = today_vn()
    if thresholds is None:
        thresholds = await load_exam_ban_thresholds(session)
    warning_percent, ban_percent = thresholds

    # 1. Fetch class details
    stmt_class = (
        select(Class)
        .options(selectinload(Class.teacher), selectinload(Class.course))
        .where(Class.id == class_id)
    )
    class_res = await session.execute(stmt_class)
    class_obj = class_res.scalar_one_or_none()
    if not class_obj:
        return None

    total_sessions = class_obj.total_sessions or 15

    # 2. Fetch enrolled students
    stmt_students = (
        select(User)
        .join(ClassStudent, ClassStudent.student_id == User.id)
        .options(selectinload(User.student_profile))
        .where(
            ClassStudent.class_id == class_id,
            User.is_active == True,
            User.deleted_at.is_(None),
        )
        .order_by(User.user_code)
    )
    students = (await session.execute(stmt_students)).scalars().all()

    # 3. Fetch schedules up to reference_date, rồi CHỈ giữ buổi đã KẾT THÚC.
    #    Buổi hôm nay chưa tới (hoặc đang học) không được tính: bản ghi điểm danh
    #    chưa tồn tại nên bước 5 sẽ coi là vắng ⇒ sinh viên bị đẩy tỷ lệ vắng và
    #    bị gắn "BỊ CẤM THI" oan cho một buổi chưa diễn ra.
    stmt_schedules = (
        select(Schedule)
        .where(Schedule.class_id == class_id, Schedule.study_date <= reference_date)
        .order_by(Schedule.study_date, Schedule.start_time)
    )
    all_schedules = (await session.execute(stmt_schedules)).scalars().all()
    # study_date/end_time lưu giờ tường địa phương (naive) nên so với giờ VN naive.
    now_local = now_vn().replace(tzinfo=None)
    schedules = [
        s for s in all_schedules
        if datetime.combine(s.study_date, s.end_time) <= now_local
    ]
    elapsed_sessions = len(schedules)

    # 4. Fetch attendances for elapsed schedules
    schedule_ids = [s.id for s in schedules]
    attendances: List[Attendance] = []
    if schedule_ids:
        stmt_att = select(Attendance).where(Attendance.schedule_id.in_(schedule_ids))
        attendances = (await session.execute(stmt_att)).scalars().all()

    att_map = {(a.schedule_id, a.student_id): a for a in attendances}

    # 5. Compute metrics for each student (Counter thay if/elif từng buổi)
    student_results = []
    for st in students:
        counts: Counter = Counter()
        for sched in schedules:
            att = att_map.get((sched.id, st.id))
            # Buổi đã kết thúc mà chưa có bản ghi = vắng không phép.
            counts[att.status if att else AttendanceStatus.ABSENT] += 1

        n_absent = counts[AttendanceStatus.ABSENT]
        n_half_absent = counts[AttendanceStatus.HALF_ABSENT]

        # V = N_ABSENT + 0.5 * N_HALF_ABSENT (EXCUSED is NOT counted)
        v = n_absent + HALF_ABSENT_WEIGHT * n_half_absent

        rate_cumulative = round((v / max(1, elapsed_sessions)) * 100.0, 2)
        rate_semester = round((v / max(1, total_sessions)) * 100.0, 2)

        is_warning = (rate_cumulative >= warning_percent)
        is_banned = (rate_semester >= ban_percent)

        if is_banned:
            status_label = "BỊ CẤM THI"
        elif is_warning:
            status_label = "CẢNH BÁO"
        else:
            status_label = "ĐỦ ĐIỀU KIỆN"

        class_name = st.student_profile.class_name if st.student_profile else ""

        student_results.append({
            "student_id": st.id,
            "student_code": st.user_code,
            "full_name": st.full_name,
            "class_name": class_name,
            "course_name": class_obj.course_name,
            "class_code": class_obj.class_code,
            "absent_count": n_absent,
            "half_absent_count": n_half_absent,
            "late_count": counts[AttendanceStatus.LATE],
            "present_count": counts[AttendanceStatus.PRESENT],
            "excused_count": counts[AttendanceStatus.EXCUSED],
            "converted_absent": v,
            "rate_cumulative": rate_cumulative,
            "rate_semester": rate_semester,
            "is_warning": is_warning,
            "is_banned": is_banned,
            "status_label": status_label,
        })

    return {
        "class_id": class_obj.id,
        "class_code": class_obj.class_code,
        "course_name": class_obj.course_name,
        "semester": class_obj.semester,
        "teacher_name": class_obj.teacher.full_name if class_obj.teacher else "",
        "total_sessions": total_sessions,
        "elapsed_sessions": elapsed_sessions,
        # Ngưỡng đang áp dụng để frontend hiển thị/tô màu khớp với backend.
        # (Admin đổi WARNING_ABSENCE_PERCENT / MAX_ABSENCE_PERCENT_EXAM_BAN ở
        # /admin/settings thì badge is_warning/is_banned và nhãn ngưỡng đổi cùng nhau.)
        "warning_percent": warning_percent,
        "ban_percent": ban_percent,
        "students": student_results,
    }


async def calculate_all_classes_exam_ban(
    session: AsyncSession,
    reference_date: Optional[date] = None,
) -> List[Dict[str, Any]]:
    """Calculate exam-ban data across all classes in the school for admin report."""
    stmt = select(Class).options(selectinload(Class.course)).order_by(Class.class_code)
    classes = (await session.execute(stmt)).scalars().all()

    # Đọc cấu hình 1 lần cho cả trường thay vì mỗi lớp 1 lần
    thresholds = await load_exam_ban_thresholds(session)

    all_rows: List[Dict[str, Any]] = []
    for c in classes:
        res = await calculate_class_exam_ban(
            session, c.id, reference_date=reference_date, thresholds=thresholds
        )
        if res and res.get("students"):
            all_rows.extend(res["students"])

    return all_rows


def _write_title_block(ws: Worksheet, title: str, subtitle: str) -> None:
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.style = "exam_title"
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:I2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = subtitle or f"Ngày xuất báo cáo: {today_vn().strftime('%d/%m/%Y')}"
    subtitle_cell.style = "exam_subtitle"
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10  # dòng phân cách


def _write_header_row(ws: Worksheet, row: int = 4) -> None:
    ws.row_dimensions[row].height = 28
    for col_idx, h in enumerate(EXAM_BAN_HEADERS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.style = "exam_header"


def _write_data_row(ws: Worksheet, row_idx: int, stt: int, item: Dict[str, Any]) -> None:
    rate_semester = item.get("rate_semester", 0.0)
    status_label = item.get("status_label", "ĐỦ ĐIỀU KIỆN")
    row_values = [
        stt,
        item.get("student_code", ""),
        item.get("full_name", ""),
        item.get("class_name", ""),
        item.get("course_name", ""),
        item.get("absent_count", 0),
        item.get("half_absent_count", 0),
        f"{rate_semester:.1f}%",
        status_label,
    ]
    ws.row_dimensions[row_idx].height = 22
    is_banned = status_label == "BỊ CẤM THI" or item.get("is_banned", False)
    is_warning = status_label == "CẢNH BÁO" or item.get("is_warning", False)
    for col_idx, val in enumerate(row_values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        centered = col_idx in (1, 2, 6, 7, 8, 9)
        if is_banned:
            if col_idx == 9:
                cell.style = "exam_banned_status"
            else:
                cell.style = "exam_banned" if centered else "exam_banned_left"
        elif is_warning:
            if col_idx == 9:
                cell.style = "exam_warning_status"
            else:
                cell.style = "exam_warning" if centered else "exam_warning_left"
        else:
            cell.style = "exam_data" if centered else "exam_data_left"


def _autofit_columns(ws: Worksheet, skip_rows: tuple[int, ...] = (1, 2, 3)) -> None:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(
            (len(str(c.value or "")) for c in col if c.row not in skip_rows),
            default=0,
        )
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)


def generate_excel_exam_ban(
    data: List[Dict[str, Any]],
    title: str,
    subtitle: str = "",
) -> bytes:
    """Generate professional Excel (.xlsx) file for exam-ban report using openpyxl."""
    wb = Workbook()
    _ensure_styles(wb)
    ws = wb.active
    ws.title = "Danh sách cấm thi"

    _write_title_block(ws, title, subtitle)
    _write_header_row(ws)
    for row_idx, item in enumerate(data, start=5):
        _write_data_row(ws, row_idx, row_idx - 4, item)
    _autofit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
