"""Bảng điểm rèn luyện theo lớp học phần.

Chính sách: điểm global (student_profiles.total_point) + delta trong lớp
(sum point_ledger theo schedule của lớp) + rank global từ classify_conduct.
Không ghi ledger mới, không đổi total_point. Thuần tính lúc đọc.
"""

from __future__ import annotations

from collections import defaultdict
from typing import Any, Dict, List, Optional

from detect_ai.modules.evaluation.public import classify_conduct

CONDUCT_EXCEL_HEADERS = [
    "STT",
    "Mã SV",
    "Họ và Tên",
    "Lớp Sinh Hoạt",
    "Điểm RL",
    "Xếp Loại",
    "Biến Động Lớp",
    "Số Mục Trong Lớp",
]


def parse_month_param(month: str | None) -> tuple[int, int] | None:
    """'YYYY-MM' -> (year, month); None/"" -> None (toàn kỳ). Raise ValueError nếu sai format."""
    if not month:
        return None
    parts = month.split("-")
    if len(parts) != 2 or len(parts[0]) != 4 or not parts[0].isdigit() or not parts[1].isdigit():
        raise ValueError(f"month phải dạng YYYY-MM, nhận: {month!r}")
    year, mon = int(parts[0]), int(parts[1])
    if not 1 <= mon <= 12:
        raise ValueError(f"month phải dạng YYYY-MM, nhận: {month!r}")
    return year, mon


def build_class_conduct_rows(
    students: List[Any],
    schedules: List[Any],
    ledger_entries: List[Any],
) -> List[Dict[str, Any]]:
    """Ghép enrolled students + ledger lớp -> rows bảng điểm.

    students: User có .student_profile (total_point).
    schedules: Schedule của lớp (đếm evaluated = schedule có ledger).
    ledger_entries: PointLedger thuộc các schedule của lớp.
    """
    delta_by_student: Dict[int, int] = defaultdict(int)
    evaluated_schedules: set[int] = set()
    entries_by_student: Dict[int, int] = defaultdict(int)
    for entry in ledger_entries:
        sid = getattr(entry, "student_id", None)
        if sid is None:
            continue
        delta_by_student[sid] += getattr(entry, "point_delta", 0) or 0
        entries_by_student[sid] += 1
        sched_id = getattr(entry, "schedule_id", None)
        if sched_id is not None:
            evaluated_schedules.add(sched_id)

    evaluated_count = len(evaluated_schedules)
    rows: List[Dict[str, Any]] = []
    for st in students:
        profile = getattr(st, "student_profile", None)
        total = getattr(profile, "total_point", None) if profile else None
        is_officer = bool(getattr(profile, "is_officer", False)) if profile else False
        total_point = total if total is not None else (100 if is_officer else 95)
        sid = getattr(st, "id", None)
        rows.append(
            {
                "student_id": sid,
                "student_code": getattr(st, "user_code", ""),
                "full_name": getattr(st, "full_name", ""),
                "class_name": getattr(profile, "class_name", "") or "",
                "total_point": total_point,
                "is_officer": is_officer,
                "conduct_rank": classify_conduct(total_point, is_officer),
                "class_delta": delta_by_student.get(sid, 0),
                "class_entries": entries_by_student.get(sid, 0),
                "evaluated_sessions": evaluated_count,
                "total_sessions": len(schedules),
            }
        )
    rows.sort(key=lambda r: (-r["total_point"], r["full_name"]))
    return rows


def generate_excel_conduct(
    rows: List[Dict[str, Any]],
    title: str,
    subtitle: str = "",
) -> bytes:
    """Xuất bảng điểm rèn luyện ra .xlsx, tái dùng style của báo cáo cấm thi."""
    from .build import _autofit_columns, _ensure_styles

    from openpyxl import Workbook

    wb = Workbook()
    _ensure_styles(wb)
    ws = wb.active
    ws.title = "Điểm rèn luyện"

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.style = "exam_title"
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:H2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = subtitle
    subtitle_cell.style = "exam_subtitle"
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10

    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(CONDUCT_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.style = "exam_header"

    for idx, r in enumerate(rows, start=1):
        row_idx = idx + 4
        delta = r.get("class_delta", 0)
        delta_str = f"{delta:+d}" if delta else "±0"
        values = [
            idx,
            r.get("student_code", ""),
            r.get("full_name", ""),
            r.get("class_name", ""),
            r.get("total_point", 0),
            r.get("conduct_rank", ""),
            delta_str,
            r.get("class_entries", 0),
        ]
        ws.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            centered = col_idx in (1, 2, 5, 6, 7, 8)
            cell.style = "exam_data" if centered else "exam_data_left"

    _autofit_columns(ws)

    import io

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def summarize_conduct(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Đếm phân bố 6 mức rank cho KPI header."""
    dist: Dict[str, int] = {}
    for r in rows:
        rank = r.get("conduct_rank", "")
        dist[rank] = dist.get(rank, 0) + 1
    return {"total": len(rows), "by_rank": dist}
