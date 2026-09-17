from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from ..models.academic import Class, ClassStudent, Schedule
from ..models.attendance_v2 import Attendance, AttendanceStatus
from ..models.user import User


async def calculate_class_exam_ban(
    session: AsyncSession,
    class_id: int,
    reference_date: Optional[date] = None,
) -> Optional[Dict[str, Any]]:
    """Calculate exam-ban metrics and status for all students in a specific class."""
    if reference_date is None:
        reference_date = date.today()

    # 1. Fetch class details
    stmt_class = (
        select(Class)
        .options(selectinload(Class.teacher))
        .where(Class.id == class_id)
    )
    class_res = await session.execute(stmt_class)
    class_obj = class_res.scalar_one_or_none()
    if not class_obj:
        return None

    total_sessions = class_obj.total_sessions or 15

    # 2. Fetch enrolled students
    stmt_students = (
        select(User)
        .join(ClassStudent, ClassStudent.student_id == User.id)
        .options(selectinload(User.student_profile))
        .where(
            ClassStudent.class_id == class_id,
            User.is_active == True,
            User.deleted_at.is_(None),
        )
        .order_by(User.user_code)
    )
    students = (await session.execute(stmt_students)).scalars().all()

    # 3. Fetch all schedules up to reference_date (elapsed sessions)
    stmt_schedules = (
        select(Schedule)
        .where(Schedule.class_id == class_id, Schedule.study_date <= reference_date)
        .order_by(Schedule.study_date, Schedule.start_time)
    )
    schedules = (await session.execute(stmt_schedules)).scalars().all()
    elapsed_sessions = len(schedules)

    # 4. Fetch attendances for elapsed schedules
    schedule_ids = [s.id for s in schedules]
    attendances: List[Attendance] = []
    if schedule_ids:
        stmt_att = select(Attendance).where(Attendance.schedule_id.in_(schedule_ids))
        attendances = (await session.execute(stmt_att)).scalars().all()

    att_map = {(a.schedule_id, a.student_id): a for a in attendances}

    # 5. Compute metrics for each student
    student_results = []
    for st in students:
        n_absent = 0
        n_half_absent = 0
        n_late = 0
        n_present = 0
        n_excused = 0

        for sched in schedules:
            att = att_map.get((sched.id, st.id))
            if att:
                if att.status == AttendanceStatus.ABSENT:
                    n_absent += 1
                elif att.status == AttendanceStatus.HALF_ABSENT:
                    n_half_absent += 1
                elif att.status == AttendanceStatus.LATE:
                    n_late += 1
                elif att.status == AttendanceStatus.PRESENT:
                    n_present += 1
                elif att.status == AttendanceStatus.EXCUSED:
                    n_excused += 1
            else:
                # If no attendance record exists for a past schedule, student was absent
                n_absent += 1

        # V = N_ABSENT + 0.5 * N_HALF_ABSENT (EXCUSED is NOT counted)
        v = n_absent + 0.5 * n_half_absent

        rate_cumulative = round((v / max(1, elapsed_sessions)) * 100.0, 2)
        rate_semester = round((v / max(1, total_sessions)) * 100.0, 2)

        is_warning = (rate_cumulative >= 15.0)
        is_banned = (rate_semester >= 20.0)

        if is_banned:
            status_label = "BỊ CẤM THI"
        elif is_warning:
            status_label = "CẢNH BÁO"
        else:
            status_label = "ĐỦ ĐIỀU KIỆN"

        class_name = st.student_profile.class_name if st.student_profile else ""

        student_results.append({
            "student_id": st.id,
            "student_code": st.user_code,
            "full_name": st.full_name,
            "class_name": class_name,
            "course_name": class_obj.course_name,
            "class_code": class_obj.class_code,
            "absent_count": n_absent,
            "half_absent_count": n_half_absent,
            "late_count": n_late,
            "present_count": n_present,
            "excused_count": n_excused,
            "converted_absent": v,
            "rate_cumulative": rate_cumulative,
            "rate_semester": rate_semester,
            "is_warning": is_warning,
            "is_banned": is_banned,
            "status_label": status_label,
        })

    return {
        "class_id": class_obj.id,
        "class_code": class_obj.class_code,
        "course_name": class_obj.course_name,
        "semester": class_obj.semester,
        "teacher_name": class_obj.teacher.full_name if class_obj.teacher else "",
        "total_sessions": total_sessions,
        "elapsed_sessions": elapsed_sessions,
        "students": student_results,
    }


async def calculate_all_classes_exam_ban(
    session: AsyncSession,
    reference_date: Optional[date] = None,
) -> List[Dict[str, Any]]:
    """Calculate exam-ban data across all classes in the school for admin report."""
    stmt = select(Class).order_by(Class.class_code)
    classes = (await session.execute(stmt)).scalars().all()

    all_rows: List[Dict[str, Any]] = []
    for c in classes:
        res = await calculate_class_exam_ban(session, c.id, reference_date=reference_date)
        if res and res.get("students"):
            all_rows.extend(res["students"])

    return all_rows


def generate_excel_exam_ban(
    data: List[Dict[str, Any]],
    title: str,
    subtitle: str = "",
) -> bytes:
    """Generate professional Excel (.xlsx) file for exam-ban report using openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách cấm thi"

    # Define color styles
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    white_font_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="4B5563")

    banned_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    banned_font = Font(name="Calibri", size=11, bold=True, color="991B1B")

    warning_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    warning_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    regular_font = Font(name="Calibri", size=11)

    thin_border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    # 1. Title row
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # 2. Subtitle row
    ws.merge_cells("A2:I2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = subtitle or f"Ngày xuất báo cáo: {date.today().strftime('%d/%m/%Y')}"
    subtitle_cell.font = subtitle_font
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Row 3 is empty separator
    ws.row_dimensions[3].height = 10

    # 3. Header row (Row 4)
    headers = [
        "STT",
        "Mã Sinh Viên",
        "Họ và Tên",
        "Lớp Sinh Hoạt",
        "Môn Học",
        "Số Buổi Vắng",
        "Số Buổi Vắng Nửa Buổi",
        "Tỷ Lệ Vắng Cả Kỳ (%)",
        "Trạng Thái",
    ]
    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = white_font_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    # 4. Data rows (Row 5+)
    for row_idx, item in enumerate(data, start=5):
        stt = row_idx - 4
        student_code = item.get("student_code", "")
        full_name = item.get("full_name", "")
        class_name = item.get("class_name", "")
        course_name = item.get("course_name", "")
        absent_count = item.get("absent_count", 0)
        half_absent_count = item.get("half_absent_count", 0)
        rate_semester = item.get("rate_semester", 0.0)
        status_label = item.get("status_label", "ĐỦ ĐIỀU KIỆN")

        row_values = [
            stt,
            student_code,
            full_name,
            class_name,
            course_name,
            absent_count,
            half_absent_count,
            f"{rate_semester:.1f}%",
            status_label,
        ]

        ws.row_dimensions[row_idx].height = 22

        is_banned = (status_label == "BỊ CẤM THI" or item.get("is_banned", False))
        is_warning = (status_label == "CẢNH BÁO" or item.get("is_warning", False))

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = cell_border
            cell.font = regular_font

            # Text alignment
            if col_idx in [1, 2, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Apply row / status highlights
            if is_banned:
                cell.fill = banned_fill
                if col_idx == 9:
                    cell.font = banned_font
            elif is_warning:
                cell.fill = warning_fill
                if col_idx == 9:
                    cell.font = warning_font

    # Auto adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip merged title rows when calculating widths
            if cell.row in [1, 2, 3]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
