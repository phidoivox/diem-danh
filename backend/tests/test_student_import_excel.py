"""Import SV từ Excel: header tự dò, pass 123456, skip trùng + báo lỗi dòng."""

import io

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import Workbook

from detect_ai.main import app  # noqa: F401
from detect_ai.modules.identity.infrastructure.models import Role, User
from detect_ai.shared.kernel.security import create_access_token, verify_password


@pytest_asyncio.fixture
async def seeded_admin(session_maker_fixture):
    async with session_maker_fixture() as session:
        session.add_all([Role(role_name="ADMIN"), Role(role_name="TEACHER"), Role(role_name="STUDENT")])
        await session.commit()
    async with session_maker_fixture() as session:
        yield session


@pytest_asyncio.fixture
async def client(make_client):
    async with make_client() as ac:
        yield ac


def make_headers():
    token = create_access_token(data={"sub": "AD001", "role": "ADMIN"})
    return {"Authorization": f"Bearer {token}"}


async def seed_admin(session_maker_fixture):
    from detect_ai.shared.kernel.security import get_password_hash

    async with session_maker_fixture() as s:
        role = [r for r in (await s.execute(__import__("sqlalchemy").select(Role))).scalars().all() if r.role_name == "ADMIN"][0]
        s.add(User(user_code="AD001", full_name="Admin", email="admin@school.edu.vn",
                   password_hash=get_password_hash("Admin@123"), role_id=role.id, is_active=True))
        await s.commit()


def xlsx_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_import_ok_auto_email_and_login(client: AsyncClient, session_maker_fixture, seeded_admin):
    await seed_admin(session_maker_fixture)
    payload = xlsx_bytes(
        ["Mã Sinh Viên", "Họ và Tên", "Email", "Lớp Sinh Hoạt", "Chuyên Ngành"],
        [("SV1001", "Nguyễn Văn A", "a@student.edu.vn", "CNTT-K24A", "Kỹ thuật phần mềm"),
         ("SV1002", "Trần Thị B", "", "CNTT-K24A", "AI")],
    )
    res = await client.post("/api/v1/admin/users/import-excel", files={"file": ("sv.xlsx", payload,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=make_headers())
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 2 and body["skipped"] == 0 and body["default_password"] == "123456"

    # Email trống -> tự sinh; pass 123456 đăng nhập được.
    login = await client.post("/api/v1/auth/login", json={"user_code_or_email": "SV1002", "password": "123456"})
    assert login.status_code == 200, login.text
    async with session_maker_fixture() as s:
        from sqlalchemy import select
        u = (await s.execute(select(User).where(User.user_code == "SV1001"))).scalar_one()
        assert u.email == "a@student.edu.vn"
        assert verify_password("123456", u.password_hash)
        assert u.student_profile.class_name == "CNTT-K24A"


@pytest.mark.asyncio
async def test_import_alias_no_accent_and_dup_skip(client: AsyncClient, session_maker_fixture, seeded_admin):
    await seed_admin(session_maker_fixture)
    first = xlsx_bytes(["mssv", "ho ten", "mail", "lop", "nganh"], [("SV2001", "Lê Văn C", "", "L1", "N1")])
    r1 = await client.post("/api/v1/admin/users/import-excel", files={"file": ("a.xlsx", first,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=make_headers())
    assert r1.json()["created"] == 1

    second = xlsx_bytes(["MSSV", "HO TEN", "MAIL", "LOP", "NGANH"],
        [("SV2001", "Trùng Mã", "trung@x.vn", "L1", "N1"),  # trùng mã
         ("SV2002", "Người Mới", "le@x.vn", "L2", "N2")])
    r2 = await client.post("/api/v1/admin/users/import-excel", files={"file": ("b.xlsx", second,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=make_headers())
    body = r2.json()
    assert body["created"] == 1 and body["skipped"] == 1
    assert body["errors"][0]["row"] == 2 and "đã tồn tại" in body["errors"][0]["message"]


@pytest.mark.asyncio
async def test_import_template_and_rejections(client: AsyncClient, session_maker_fixture, seeded_admin):
    await seed_admin(session_maker_fixture)
    t = await client.get("/api/v1/admin/users/import-excel/template", headers=make_headers())
    assert t.status_code == 200
    assert "spreadsheetml" in t.headers["content-type"]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(t.content)).active
    assert [ws.cell(row=1, column=c).value for c in range(1, 6)] == ["Mã Sinh Viên", "Họ và Tên", "Email", "Lớp Sinh Hoạt", "Chuyên Ngành"]

    bad_ext = await client.post("/api/v1/admin/users/import-excel", files={"file": ("sv.txt", b"hello",
        "text/plain")}, headers=make_headers())
    assert bad_ext.status_code == 400

    no_header = xlsx_bytes(["Cột A", "Cột B"], [("x", "y")])
    r = await client.post("/api/v1/admin/users/import-excel", files={"file": ("sv.xlsx", no_header,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=make_headers())
    assert r.status_code == 400 and "tiêu đề" in r.json()["detail"]
